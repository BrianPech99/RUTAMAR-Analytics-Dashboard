import csv
from io import BytesIO
from datetime import datetime

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
import pandas as pd
from pydantic import BaseModel, Field

from app.services.rutamar_data import CAPTURE_FILE, capture_options, passengers

router = APIRouter(tags=["capture"])


class CaptureRecord(BaseModel):
    ruta_csv: str
    unidad: str
    sentido: str
    paradero: str
    horario_programado: str
    ascensos: int = Field(ge=0)
    descensos: int = Field(ge=0)
    fecha_hora_real: datetime
    observaciones: str = Field(default="", max_length=500)


@router.get("/capture/options")
def options() -> dict:
    return {"rutas": capture_options()}


@router.get("/capture/export")
def export_captures() -> StreamingResponse:
    try:
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise HTTPException(
            status_code=503,
            detail="La exportación a Excel requiere openpyxl. Instala las dependencias de backend con: pip install -r requirements.txt",
        ) from error

    data = passengers()
    columns = ["Edicion", "Fecha", "Dia", "Ruta", "Unidad", "Sentido", "Paradero", "Hora Programada", "Hora Real de Llegada", "Ascensos (Suben)", "Descensos (Bajan)", "Observaciones"]
    output = BytesIO()
    try:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            data.reindex(columns=columns).to_excel(writer, sheet_name="Registros", index=False)
            sheet = writer.sheets["Registros"]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            for index, column in enumerate(sheet.columns, start=1):
                width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
                sheet.column_dimensions[get_column_letter(index)].width = width
    except Exception as error:
        raise HTTPException(status_code=500, detail="No fue posible generar el archivo Excel.") from error
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=RUTAMAR-Registros.xlsx"})


@router.post("/capture")
def save_capture(record: CaptureRecord) -> dict:
    valid_routes = {route["ruta_csv"] for route in capture_options()}
    if record.ruta_csv not in valid_routes or record.sentido not in {"Ida", "Regreso"}:
        raise HTTPException(status_code=422, detail="La ruta o el sentido no es válido.")

    timestamp = record.fecha_hora_real
    headers = ["Edicion", "Fecha", "Dia", "Ruta", "Unidad", "Sentido", "Paradero", "Hora Programada", "Hora Real de Llegada", "Ascensos (Suben)", "Descensos (Bajan)", "Observaciones"]
    row = {
        "Edicion": "Captura manual",
        "Fecha": timestamp.strftime("%Y-%m-%d"),
        "Dia": timestamp.strftime("%A"),
        "Ruta": record.ruta_csv,
        "Unidad": record.unidad,
        "Sentido": record.sentido,
        "Paradero": record.paradero,
        "Hora Programada": record.horario_programado,
        "Hora Real de Llegada": timestamp.strftime("%H:%M"),
        "Ascensos (Suben)": record.ascensos,
        "Descensos (Bajan)": record.descensos,
        "Observaciones": record.observaciones.strip(),
    }
    is_new_file = not CAPTURE_FILE.exists()
    with CAPTURE_FILE.open("a", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        if is_new_file:
            writer.writeheader()
        writer.writerow(row)
    passengers.cache_clear()
    return {"ok": True, "message": "Registro guardado e incorporado al dashboard."}
